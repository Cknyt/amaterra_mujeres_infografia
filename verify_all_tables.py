import openpyxl
import json
import subprocess
import re

wb = openpyxl.load_workbook('20260925_DatosExplotaciones.xlsx', data_only=True)
ws_no = wb['No_ATP']
ws_atp = wb['ATP']

with open('reporte_diputacion_agro_2.html', 'r', encoding='utf-8') as f:
    html = f.read()

m = re.search(r'const DATA_EXCEL = (\{.+?\});\s*(?:const|function|let|var)', html, re.DOTALL)
js_code = m.group(1)
with open('temp_data.js', 'w', encoding='utf-8') as f:
    f.write('console.log(JSON.stringify(' + js_code + '));')

res = subprocess.run(['node', 'temp_data.js'], capture_output=True, text=True, encoding='utf-8')
data = json.loads(res.stdout)

print("="*60)
print("VERIFYING DATA_EXCEL AGAINST EXCEL FILE")
print("="*60)

# 1. ATP SUBSECTORES
print("\n[1] ATP SUBSECTORES:")
excel_atp_sub = []
for r in range(12, 27):
    name = str(ws_atp.cell(r, 4).value).strip()
    h = int(ws_atp.cell(r, 5).value or 0)
    m_val = int(ws_atp.cell(r, 6).value or 0)
    pct = round(float(ws_atp.cell(r, 9).value or 0), 2)
    excel_atp_sub.append({'name': name, 'h': h, 'm': m_val, 'pct': pct})

html_atp_sub = data['ATP']['subsectores']
print(f"Excel count: {len(excel_atp_sub)}, HTML count: {len(html_atp_sub)}")
for i, item in enumerate(html_atp_sub):
    e = excel_atp_sub[i]
    diff = []
    if e['h'] != item['h']: diff.append(f"H: {e['h']} vs {item['h']}")
    if e['m'] != item['m']: diff.append(f"M: {e['m']} vs {item['m']}")
    if abs(e['pct'] - item['pct_m_de_190']) > 0.05: diff.append(f"pct: {e['pct']} vs {item['pct_m_de_190']}")
    if diff:
        print(f"  DIFF at index {i} ({item['nombre']}): {', '.join(diff)}")
print("ATP subsectors verified.")

# 2. NO_ATP SUBSECTORES
print("\n[2] NO_ATP SUBSECTORES:")
excel_no_sub = []
for r in range(30, 48):
    name = str(ws_no.cell(r, 4).value).strip()
    h = int(ws_no.cell(r, 5).value or 0)
    m_val = int(ws_no.cell(r, 6).value or 0)
    pct = round(float(ws_no.cell(r, 8).value or 0), 2)
    excel_no_sub.append({'name': name, 'h': h, 'm': m_val, 'pct': pct})

html_no_sub = data['No_ATP']['subsectores']
print(f"Excel count: {len(excel_no_sub)}, HTML count: {len(html_no_sub)}")
for i, item in enumerate(html_no_sub):
    e = excel_no_sub[i]
    diff = []
    if e['h'] != item['h']: diff.append(f"H: {e['h']} vs {item['h']}")
    if e['m'] != item['m']: diff.append(f"M: {e['m']} vs {item['m']}")
    if abs(e['pct'] - item['pct_m_de_1004']) > 0.05: diff.append(f"pct: {e['pct']} vs {item['pct_m_de_1004']}")
    if diff:
        print(f"  DIFF at index {i} ({item['nombre']}): {', '.join(diff)}")
print("NO_ATP subsectors verified.")

# 3. SUPERFICIE ATP
print("\n[3] ATP SUPERFICIE:")
excel_atp_sup = []
for r in range(33, 37):
    name = str(ws_atp.cell(r, 4).value).strip()
    h = int(ws_atp.cell(r, 5).value or 0)
    m_val = int(ws_atp.cell(r, 7).value or 0)
    pct = round(float(ws_atp.cell(r, 9).value or 0), 2)
    excel_atp_sup.append({'name': name, 'h': h, 'm': m_val, 'pct': pct})

for i, item in enumerate(data['ATP']['superficie']):
    e = excel_atp_sup[i]
    diff = []
    if e['h'] != item['h']: diff.append(f"H: {e['h']} vs {item['h']}")
    if e['m'] != item['m']: diff.append(f"M: {e['m']} vs {item['m']}")
    if abs(e['pct'] - item['pct_m_de_190']) > 0.05: diff.append(f"pct: {e['pct']} vs {item['pct_m_de_190']}")
    if diff:
        print(f"  DIFF in ATP superficie: {diff}")
print("ATP superficie verified.")

# 4. SUPERFICIE NO ATP
print("\n[4] NO_ATP SUPERFICIE:")
excel_no_sup = []
for r in range(57, 61):
    name = str(ws_no.cell(r, 4).value).strip()
    h = int(ws_no.cell(r, 5).value or 0)
    m_val = int(ws_no.cell(r, 7).value or 0)
    pct = round(float(ws_no.cell(r, 8).value or 0), 2)
    excel_no_sup.append({'name': name, 'h': h, 'm': m_val, 'pct': pct})

for i, item in enumerate(data['No_ATP']['superficie']):
    e = excel_no_sup[i]
    diff = []
    if e['h'] != item['h']: diff.append(f"H: {e['h']} vs {item['h']}")
    if e['m'] != item['m']: diff.append(f"M: {e['m']} vs {item['m']}")
    if abs(e['pct'] - item['pct_m_de_1004']) > 0.05: diff.append(f"pct: {e['pct']} vs {item['pct_m_de_1004']}")
    if diff:
        print(f"  DIFF in No_ATP superficie: {diff}")
print("NO_ATP superficie verified.")

# 5. COMARCAS ATP
print("\n[5] ATP COMARCAS:")
excel_atp_com = []
for r in range(45, 51):
    name = str(ws_atp.cell(r, 4).value).strip()
    h = int(ws_atp.cell(r, 5).value or 0)
    m_val = int(ws_atp.cell(r, 7).value or 0)
    pct = round(float(ws_atp.cell(r, 9).value or 0), 2)
    excel_atp_com.append({'name': name, 'h': h, 'm': m_val, 'pct': pct})

for i, item in enumerate(data['ATP']['comarca']):
    e = excel_atp_com[i]
    diff = []
    if e['h'] != item['h']: diff.append(f"H: {e['h']} vs {item['h']}")
    if e['m'] != item['m']: diff.append(f"M: {e['m']} vs {item['m']}")
    if abs(e['pct'] - item['pct_m_de_190']) > 0.05: diff.append(f"pct: {e['pct']} vs {item['pct_m_de_190']}")
    if diff:
        print(f"  DIFF in ATP comarca {item['comarca']}: {diff}")
print("ATP comarcas verified.")

# 6. COMARCAS NO ATP
print("\n[6] NO_ATP COMARCAS:")
excel_no_com = []
for r in range(71, 77):
    name = str(ws_no.cell(r, 4).value).strip()
    h = int(ws_no.cell(r, 5).value or 0)
    m_val = int(ws_no.cell(r, 7).value or 0)
    pct = round(float(ws_no.cell(r, 8).value or 0), 2)
    excel_no_com.append({'name': name, 'h': h, 'm': m_val, 'pct': pct})

for i, item in enumerate(data['No_ATP']['comarca']):
    e = excel_no_com[i]
    diff = []
    if e['h'] != item['h']: diff.append(f"H: {e['h']} vs {item['h']}")
    if e['m'] != item['m']: diff.append(f"M: {e['m']} vs {item['m']}")
    if abs(e['pct'] - item['pct_m_de_1004']) > 0.05: diff.append(f"pct: {e['pct']} vs {item['pct_m_de_1004']}")
    if diff:
        print(f"  DIFF in No_ATP comarca {item['comarca']}: {diff}")
print("NO_ATP comarcas verified.")

# 7. ECOLOGICO ATP
print("\n[7] ATP ECOLOGICO:")
h_eco = int(ws_atp.cell(40, 5).value) # 32
m_eco = int(ws_atp.cell(40, 7).value) # 22
m_pct = round(float(ws_atp.cell(40, 8).value), 2) # 11.58
item = data['ATP']['ecologico'][0]
print(f"Excel: H={h_eco}, M={m_eco}, M%={m_pct} | HTML: H={item['h']}, M={item['m']}, M%={item['pct_m_de_190']}")

# 8. EDAD ATP
print("\n[8] ATP EDAD:")
excel_atp_edad = []
for r in range(30, 33):
    name = str(ws_atp.cell(r, 16).value).strip()
    h = int(ws_atp.cell(r, 17).value)
    m_val = int(ws_atp.cell(r, 18).value)
    hp = round(float(ws_atp.cell(r, 19).value), 2)
    mp = round(float(ws_atp.cell(r, 20).value), 2)
    excel_atp_edad.append({'tramo': name, 'h': h, 'm': m_val, 'h_pct': hp, 'm_pct': mp})

for i, item in enumerate(data['ATP']['edad']['tramos']):
    e = excel_atp_edad[i]
    diff = []
    if e['h'] != item['h']: diff.append(f"H: {e['h']} vs {item['h']}")
    if e['m'] != item['m']: diff.append(f"M: {e['m']} vs {item['m']}")
    if abs(e['h_pct'] - item['h_pct']) > 0.05: diff.append(f"hpct: {e['h_pct']} vs {item['h_pct']}")
    if abs(e['m_pct'] - item['m_pct_sobre_m']) > 0.05: diff.append(f"mpct: {e['m_pct']} vs {item['m_pct_sobre_m']}")
    if diff:
        print(f"  DIFF in ATP edad {item['tramo']}: {diff}")
print("ATP edad verified.")

# 9. EDAD NO ATP
print("\n[9] NO_ATP EDAD:")
excel_no_edad = []
for r in range(20, 23):
    name = str(ws_no.cell(r, 3).value).strip()
    h = int(ws_no.cell(r, 4).value)
    m_val = int(ws_no.cell(r, 5).value)
    hp = round(float(ws_no.cell(r, 6).value), 2)
    mp = round(float(ws_no.cell(r, 7).value), 2)
    excel_no_edad.append({'tramo': name, 'h': h, 'm': m_val, 'h_pct': hp, 'm_pct': mp})

for i, item in enumerate(data['No_ATP']['edad']['tramos']):
    e = excel_no_edad[i]
    diff = []
    if e['h'] != item['h']: diff.append(f"H: {e['h']} vs {item['h']}")
    if e['m'] != item['m']: diff.append(f"M: {e['m']} vs {item['m']}")
    if abs(e['h_pct'] - item['h_pct']) > 0.05: diff.append(f"hpct: {e['h_pct']} vs {item['h_pct']}")
    if abs(e['m_pct'] - item['m_pct_sobre_m']) > 0.05: diff.append(f"mpct: {e['m_pct']} vs {item['m_pct_sobre_m']}")
    if diff:
        print(f"  DIFF in No_ATP edad {item['tramo']}: {diff}")
print("NO_ATP edad verified.")
