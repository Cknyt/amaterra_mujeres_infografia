# audit_excel_tables.py
import openpyxl
import json

wb = openpyxl.load_workbook('20260925_DatosExplotaciones.xlsx', data_only=True)

print("="*60)
print("AUDITORIA COMPLETA DE HOJA: No_ATP")
print("="*60)
ws_no = wb['No_ATP']
for r in range(1, ws_no.max_row + 1):
    vals = [ws_no.cell(r, c).value for c in range(1, 15)]
    # Filter none at end
    while vals and vals[-1] is None:
        vals.pop()
    if vals:
        # short print
        row_clean = [v if not (isinstance(v, str) and len(v) > 50) else (v[:40] + '...') for v in vals]
        print(f"R{r:2d}: {row_clean}")

print("\n" + "="*60)
print("AUDITORIA COMPLETA DE HOJA: ATP")
print("="*60)
ws_atp = wb['ATP']
for r in range(1, ws_atp.max_row + 1):
    vals = [ws_atp.cell(r, c).value for c in range(1, 35)]
    while vals and vals[-1] is None:
        vals.pop()
    if vals:
        row_clean = [v if not (isinstance(v, str) and len(v) > 50) else (v[:40] + '...') for v in vals]
        print(f"R{r:2d}: {row_clean}")
