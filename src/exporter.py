"""Motor de exportación y reporting automatizado (src/exporter.py).

Responsabilidades principales:
1. Exportar figuras interactivas de Plotly a imágenes vectoriales (.svg) y rasterizadas
   en alta resolución (PNG a 300 DPI usando Kaleido).
2. Generar libros consolidados de Microsoft Excel (.xlsx) con formateo corporativo
   profesional en openpyxl (cabeceras #1F4E38, formatos numéricos y autoajuste de columnas).
3. Construir presentaciones ejecutivas en Microsoft PowerPoint (.pptx, 16:9 widescreen)
   100% editables mediante python-pptx, combinando tarjetas KPI, imágenes de gráficos
   en alta resolución y tablas nativas de PowerPoint con celdas y textos editables.
"""

from datetime import datetime
from io import BytesIO
import os
from pathlib import Path
import sys
from typing import Any, Dict, List, Optional, Tuple, Union

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import plotly.graph_objects as go
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Asegurar que el directorio raíz esté en sys.path
_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from src.config import CHARTS_DIR, EXPORTS_DIR, TABLES_DIR

# Paleta Corporativa Institucional
PALETTE_PPTX = {
    "primary": RGBColor(31, 78, 56),       # Verde bosque institucional (#1F4E38)
    "primary_dark": RGBColor(27, 67, 50),  # Verde muy oscuro (#1B4332)
    "secondary_m": RGBColor(128, 90, 213), # Morado brecha (#805AD5)
    "secondary_h": RGBColor(43, 108, 176), # Azul grafito (#2B6CB0)
    "accent_red": RGBColor(197, 48, 48),   # Rojo brecha negativa (#C53030)
    "neutral_dark": RGBColor(45, 55, 72),  # Gris texto (#2D3748)
    "neutral_muted": RGBColor(113, 128, 150), # Gris anotación (#718096)
    "bg_card": RGBColor(248, 249, 250),    # Fondo tarjeta (#F8F9FA)
    "bg_zebra": RGBColor(245, 247, 250),   # Alternancia de tabla
    "white": RGBColor(255, 255, 255),
    "border_card": RGBColor(226, 232, 240),# Borde gris claro (#E2E8F0)
}


# ==============================================================================
# 1. EXPORTACIÓN DE GRÁFICOS (PNG 300 DPI & SVG)
# ==============================================================================

def export_chart(
    fig: go.Figure,
    filename_base: str,
    output_dir: Union[str, Path] = "exports/charts",
) -> Dict[str, str]:
    """Exporta un gráfico Plotly a imagen rasterizada (PNG a 300 DPI) y vectorial (SVG).

    Args:
        fig: Figura interactiva de Plotly.
        filename_base: Nombre base del archivo sin extensión.
        output_dir: Directorio de destino.

    Returns:
        Diccionario con rutas: {"png": path_png, "svg": path_svg}.
    """
    out_dir = Path(output_dir)
    os.makedirs(out_dir, exist_ok=True)

    # Limpiar extensión previa si se incluyó
    base = filename_base.replace(".png", "").replace(".svg", "")
    png_path = out_dir / f"{base}.png"
    svg_path = out_dir / f"{base}.svg"

    # Exportar PNG a alta resolución (1600x900 con escala 3 = 4800x2700 px ~ 300 DPI)
    try:
        png_bytes = fig.to_image(format="png", width=1600, height=900, scale=3)
        with open(png_path, "wb") as f:
            f.write(png_bytes)
    except Exception as e:
        print(f"[Advertencia] Error al exportar PNG con Kaleido para '{base}': {e}", file=sys.stderr)

    # Exportar SVG vectorial
    try:
        svg_bytes = fig.to_image(format="svg")
        with open(svg_path, "wb") as f:
            f.write(svg_bytes)
    except Exception as e:
        print(f"[Advertencia] Error al exportar SVG con Kaleido para '{base}': {e}", file=sys.stderr)

    return {"png": str(png_path), "svg": str(svg_path)}


# Alias de compatibilidad
def export_chart_to_file(fig: go.Figure, filename: str, **kwargs) -> Path:
    base = filename.replace(".png", "").replace(".svg", "")
    res = export_chart(fig, base, output_dir=CHARTS_DIR)
    return Path(res["png"] if filename.endswith(".png") else res.get("svg", res["png"]))


# ==============================================================================
# 2. EXPORTACIÓN DE TABLAS EXCEL CON OPENPYXL
# ==============================================================================

def export_tables_to_excel(
    tables_dict: Dict[str, pd.DataFrame],
    output_filepath: Union[str, Path] = "exports/tables/tablas_resumen.xlsx",
) -> Path:
    """Exporta un conjunto de DataFrames a un libro Excel consolidado con formato corporativo openpyxl.

    Aplica:
    - Cabeceras en negrita con fondo verde institucional (#1F4E38) y texto blanco.
    - Formatos numéricos automáticos (#,##0 para enteros, 0.0 para decimales, 0.0% para porcentajes).
    - Autoajuste inteligente del ancho de columnas para evitar textos truncados.
    """
    out_path = Path(output_filepath)
    os.makedirs(out_path.parent, exist_ok=True)

    wb = openpyxl.Workbook()
    # Eliminar hoja por defecto inicial
    wb.remove(wb.active)

    # Estilos openpyxl
    header_fill = PatternFill(start_color="1F4E38", end_color="1F4E38", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Segoe UI", size=10)
    data_zebra_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    for sheet_name, df in tables_dict.items():
        # Limitar longitud de hoja a 31 caracteres (límite Excel)
        safe_name = str(sheet_name).replace(":", "").replace("/", "")[:31]
        ws = wb.create_sheet(title=safe_name)
        ws.views.sheetView[0].showGridLines = True

        # 1. Escribir fila de cabecera
        headers = list(df.columns)
        ws.append(headers)

        for col_num, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
        ws.row_dimensions[1].height = 26

        # 2. Escribir filas de datos con formateo
        for row_idx, row_values in enumerate(df.values, start=2):
            ws.append(list(row_values))
            ws.row_dimensions[row_idx].height = 20
            is_even_row = (row_idx % 2 == 0)

            for col_num, (col_name, val) in enumerate(zip(headers, row_values), start=1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = data_font
                cell.border = thin_border
                if is_even_row:
                    cell.fill = data_zebra_fill

                col_lower = str(col_name).lower()

                # Detección y formato numérico inteligente
                if isinstance(val, (int, np.integer)):
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif isinstance(val, (float, np.floating)):
                    if any(k in col_lower for k in ["%", "feminiz", "porcent", "cuota", "ratio"]):
                        cell.number_format = '0.0"%"'
                    else:
                        cell.number_format = "#,##0.0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # 3. Autoajuste del ancho de columnas
        for col in ws.columns:
            max_len = 0
            for cell in col:
                cell_val = str(cell.value or "")
                max_len = max(max_len, len(cell_val))
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(str(out_path))
    return out_path


# Alias de compatibilidad
def export_consolidated_tables(tables_dict: Dict[str, pd.DataFrame], filename: str = "tablas_resumen.xlsx") -> Path:
    out_path = TABLES_DIR / filename
    return export_tables_to_excel(tables_dict, output_filepath=out_path)


# ==============================================================================
# 3. CONSTRUCTOR DE PRESENTACIÓN POWERPOINT (PowerPointReportBuilder)
# ==============================================================================

class PowerPointReportBuilder:
    """Orquesta la creación de presentaciones ejecutivas 16:9 en PowerPoint 100% editables."""

    def __init__(
        self,
        title: str = "Estudio de Explotaciones Agrarias Comerciales de Bizkaia: Brecha de Género y Dimensión Productiva",
        author: str = "Plataforma de Analítica AMATERRA Mujeres",
    ):
        self.prs = Presentation()
        # Widescreen 16:9 (13.333 x 7.5 pulgadas)
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.title = title
        self.author = author

    # --------------------------------------------------------------------------
    # MÉTODOS DE APOYO Y COMPONENTES VISUALES
    # --------------------------------------------------------------------------

    def add_header(self, slide, title: str, category_tag: str):
        """Inserta encabezado estandarizado con categoría superior y título principal."""
        # Barra lateral decorativa institucional
        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.45), Inches(0.12), Inches(0.85)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = PALETTE_PPTX["primary"]
        accent.line.fill.background()

        # Cuadro de texto para tag y título
        tx = slide.shapes.add_textbox(Inches(1.05), Inches(0.38), Inches(11.5), Inches(0.95))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_tag = tf.paragraphs[0]
        p_tag.text = category_tag.upper()
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.name = "Segoe UI"
        p_tag.font.color.rgb = PALETTE_PPTX["secondary_m"]

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.name = "Segoe UI"
        p_title.font.color.rgb = PALETTE_PPTX["neutral_dark"]
        p_title.space_before = Pt(3)

    def add_kpi_cards(
        self,
        slide,
        kpi_list: List[Dict[str, str]],
        left: Inches = Inches(0.8),
        top: Inches = Inches(4.8),
        width: Inches = Inches(11.733),
        height: Inches = Inches(1.8),
    ):
        """Dibuja un conjunto de tarjetas de métricas cuantitativas clave."""
        num_cards = max(1, len(kpi_list))
        card_w = (width - Inches(0.25 * (num_cards - 1))) / num_cards

        for i, card_info in enumerate(kpi_list):
            card_left = left + i * (card_w + Inches(0.25))

            # Contenedor de tarjeta
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, card_left, top, card_w, height
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = PALETTE_PPTX["bg_card"]
            shape.line.color.rgb = PALETTE_PPTX["border_card"]
            shape.line.width = Pt(1.5)

            tf = shape.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE

            # Valor principal (número destacado)
            p_val = tf.paragraphs[0]
            p_val.text = str(card_info.get("value", ""))
            p_val.font.size = Pt(28)
            p_val.font.bold = True
            p_val.font.name = "Segoe UI"
            val_color = card_info.get("color", "primary")
            p_val.font.color.rgb = PALETTE_PPTX.get(val_color, PALETTE_PPTX["primary"])
            p_val.alignment = PP_ALIGN.CENTER

            # Etiqueta
            p_lbl = tf.add_paragraph()
            p_lbl.text = str(card_info.get("label", "")).upper()
            p_lbl.font.size = Pt(10)
            p_lbl.font.bold = True
            p_lbl.font.name = "Segoe UI"
            p_lbl.font.color.rgb = PALETTE_PPTX["neutral_muted"]
            p_lbl.alignment = PP_ALIGN.CENTER
            p_lbl.space_before = Pt(4)

            # Subtítulo adicional si existe
            if "sub" in card_info:
                p_sub = tf.add_paragraph()
                p_sub.text = str(card_info["sub"])
                p_sub.font.size = Pt(9)
                p_sub.font.name = "Segoe UI"
                p_sub.font.color.rgb = PALETTE_PPTX["secondary_m"]
                p_sub.alignment = PP_ALIGN.CENTER
                p_sub.space_before = Pt(2)

    def add_native_table(
        self,
        slide,
        df: pd.DataFrame,
        left: Inches,
        top: Inches,
        width: Inches,
        height: Inches,
    ):
        """Crea una tabla nativa editable de PowerPoint formateada profesionalmente."""
        rows, cols = len(df) + 1, len(df.columns)
        table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
        tbl = table_shape.table

        # 1. Cabecera
        for j, col_name in enumerate(df.columns):
            cell = tbl.cell(0, j)
            cell.text = str(col_name)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PALETTE_PPTX["primary"]
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            for p in cell.text_frame.paragraphs:
                p.font.bold = True
                p.font.size = Pt(10)
                p.font.name = "Segoe UI"
                p.font.color.rgb = PALETTE_PPTX["white"]
                p.alignment = PP_ALIGN.CENTER

        # 2. Filas de datos
        for i, row in enumerate(df.values):
            is_zebra = (i % 2 == 1)
            for j, val in enumerate(row):
                cell = tbl.cell(i + 1, j)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = PALETTE_PPTX["bg_zebra"] if is_zebra else PALETTE_PPTX["white"]

                # Formatear el contenido de la celda
                col_name = str(df.columns[j]).lower()
                if pd.isna(val):
                    display_text = "-"
                    align = PP_ALIGN.CENTER
                elif isinstance(val, (int, np.integer)):
                    display_text = f"{val:,}"
                    align = PP_ALIGN.RIGHT
                elif isinstance(val, (float, np.floating)):
                    if any(k in col_name for k in ["%", "feminiz", "porcent"]):
                        display_text = f"{val:.1f}%"
                    else:
                        display_text = f"{val:.1f}"
                    align = PP_ALIGN.RIGHT
                else:
                    display_text = str(val)
                    align = PP_ALIGN.LEFT

                cell.text = display_text
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(9.5)
                    p.font.name = "Segoe UI"
                    p.font.color.rgb = PALETTE_PPTX["neutral_dark"]
                    p.alignment = align

    def add_chart_image(
        self,
        slide,
        image_path: str,
        left: Inches,
        top: Inches,
        width: Inches,
        height: Optional[Inches] = None,
    ):
        """Inserta una imagen de gráfico generada en alta resolución manteniendo proporciones."""
        if Path(image_path).exists():
            if height:
                slide.shapes.add_picture(image_path, left, top, width=width, height=height)
            else:
                slide.shapes.add_picture(image_path, left, top, width=width)
        else:
            # Fallback si no existe la imagen
            tx = slide.shapes.add_textbox(left, top, width, Inches(2.0))
            p = tx.text_frame.paragraphs[0]
            p.text = f"[Imagen no disponible: {Path(image_path).name}]"
            p.font.color.rgb = PALETTE_PPTX["accent_red"]

    def add_kpis_slide(
        self,
        slide_title: str,
        kpis: Dict[str, Any],
        category_tag: str = "INDICADORES",
    ):
        """Añade una diapositiva con tarjetas visuales de KPIs."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_header(slide, slide_title, category_tag)
        kpi_list = [{"label": str(k), "value": str(v), "color": "primary"} for k, v in kpis.items()]
        self.add_kpi_cards(slide, kpi_list, left=Inches(0.8), top=Inches(2.5), width=Inches(11.733), height=Inches(2.0))

    def add_chart_slide(
        self,
        slide_title: str,
        fig_or_path: Union[go.Figure, str, Path],
        notes: Optional[str] = None,
        category_tag: str = "VISUALIZACIÓN",
    ):
        """Añade una diapositiva con un gráfico Plotly o imagen estática."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_header(slide, slide_title, category_tag)

        if isinstance(fig_or_path, (str, Path)):
            img_path = str(fig_or_path)
        else:
            chart_res = export_chart(fig_or_path, "temp_chart_slide")
            img_path = chart_res["png"]

        chart_width = Inches(7.2 if notes else 11.5)
        self.add_chart_image(slide, img_path, left=Inches(0.8), top=Inches(1.6), width=chart_width)

        if notes:
            tx = slide.shapes.add_textbox(Inches(8.3), Inches(1.8), Inches(4.2), Inches(4.5))
            tf = tx.text_frame
            tf.word_wrap = True
            p_h = tf.paragraphs[0]
            p_h.text = "Conclusiones Clave:"
            p_h.font.bold = True
            p_h.font.size = Pt(11)
            p_h.font.name = "Segoe UI"
            p_h.font.color.rgb = PALETTE_PPTX["primary"]

            p_t = tf.add_paragraph()
            p_t.text = notes
            p_t.font.size = Pt(10)
            p_t.font.name = "Segoe UI"
            p_t.font.color.rgb = PALETTE_PPTX["neutral_dark"]
            p_t.space_before = Pt(6)

    # --------------------------------------------------------------------------
    # CONSTRUCCIÓN DE LAS 6 DIAPOSITIVAS CLAVE
    # --------------------------------------------------------------------------

    def build_slide_1_cover(self, kpis: Dict[str, Any]):
        """Slide 1: Portada ejecutiva con métricas macro."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

        # Banda lateral verde
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.45), Inches(7.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PALETTE_PPTX["primary"]
        bar.line.fill.background()

        # Texto del título
        tx = slide.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.3), Inches(3.0))
        tf = tx.text_frame
        tf.word_wrap = True

        p_tag = tf.paragraphs[0]
        p_tag.text = "INFORME EJECUTIVO Y ANÁLISIS DE GÉNERO"
        p_tag.font.size = Pt(12)
        p_tag.font.bold = True
        p_tag.font.name = "Segoe UI"
        p_tag.font.color.rgb = PALETTE_PPTX["secondary_m"]

        p_title = tf.add_paragraph()
        p_title.text = self.title
        p_title.font.size = Pt(32)
        p_title.font.bold = True
        p_title.font.name = "Segoe UI"
        p_title.font.color.rgb = PALETTE_PPTX["primary_dark"]
        p_title.space_before = Pt(8)

        p_sub = tf.add_paragraph()
        p_sub.text = f"{self.author} · Fecha: {datetime.now().strftime('%d/%m/%Y')}"
        p_sub.font.size = Pt(13)
        p_sub.font.name = "Segoe UI"
        p_sub.font.color.rgb = PALETTE_PPTX["neutral_muted"]
        p_sub.space_before = Pt(16)

        # 4 Tarjetas de Macro KPIs
        atp_data = kpis.get("atp", {})
        kpi_cards_data = [
            {"label": "Total Mercado", "value": f"{kpis.get('total_comerciales', 3323):,}", "sub": "100% Censo Comercial", "color": "primary"},
            {"label": "Feminización", "value": f"{kpis.get('pct_feminizacion_global', 35.9):.1f}%", "sub": f"{kpis.get('total_mujeres', 1194):,} titulares", "color": "secondary_m"},
            {"label": "Brecha Superficie (ATP)", "value": f"-{atp_data.get('brecha_superficie_pct', 25.5):.1f}%", "sub": f"{atp_data.get('superficie_media_hombres', 32.6):.1f} vs {atp_data.get('superficie_media_mujeres', 24.2):.1f} ha", "color": "accent_red"},
            {"label": "Brecha UTAs (ATP)", "value": f"-{atp_data.get('brecha_utas_pct', 32.4):.1f}%", "sub": f"{atp_data.get('utas_media_hombres', 2.1):.2f} vs {atp_data.get('utas_media_mujeres', 1.4):.2f} UTAs", "color": "accent_red"},
        ]
        self.add_kpi_cards(slide, kpi_cards_data, left=Inches(1.2), top=Inches(4.6), width=Inches(11.3), height=Inches(2.0))

    def build_slide_2_macro_funnel(self, sankey_img_path: str, kpis: Dict[str, Any]):
        """Slide 2: Macro-Embudo comercial (Sankey + Tabla nativa ATP/No ATP)."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_header(slide, "Estructura del Censo Comercial: Calificación ATP y Género", "Nivel 1: Macro-Embudo")

        # Imagen Sankey
        self.add_chart_image(slide, sankey_img_path, left=Inches(0.8), top=Inches(1.6), width=Inches(7.2))

        # Tabla nativa ATP vs No ATP
        atp = kpis.get("atp", {})
        no_atp = kpis.get("no_atp", {})
        df_macro = pd.DataFrame([
            {"Segmento": "ATP", "Hombres": atp.get("hombres", 339), "Mujeres": atp.get("mujeres", 190), "Total": atp.get("total", 529), "% Fem": atp.get("pct_feminizacion", 35.9)},
            {"Segmento": "No ATP", "Hombres": no_atp.get("hombres", 1790), "Mujeres": no_atp.get("mujeres", 1004), "Total": no_atp.get("total", 2794), "% Fem": no_atp.get("pct_feminizacion", 35.9)},
            {"Segmento": "Total Mercado", "Hombres": kpis.get("total_hombres", 2129), "Mujeres": kpis.get("total_mujeres", 1194), "Total": kpis.get("total_comerciales", 3323), "% Fem": kpis.get("pct_feminizacion_global", 35.9)},
        ])
        self.add_native_table(slide, df_macro, left=Inches(8.3), top=Inches(1.8), width=Inches(4.3), height=Inches(2.2))

        # Cuadro de notas explicativo
        tx = slide.shapes.add_textbox(Inches(8.3), Inches(4.3), Inches(4.3), Inches(2.3))
        tf = tx.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = "Hallazgos Clave:"
        p_h.font.bold = True
        p_h.font.size = Pt(11)
        p_h.font.color.rgb = PALETTE_PPTX["primary"]

        p_t = tf.add_paragraph()
        p_t.text = (
            "• El 84,1% del censo de mercado (2.794 explotaciones) corresponde a explotaciones No ATP.\n"
            "• La representatividad femenina se mantiene constante en el 35,9% tanto en el segmento profesional (ATP) como en el complementario (No ATP).\n"
            "• Solo 190 mujeres ostentan la titularidad de explotaciones agrarias prioritarias en el territorio."
        )
        p_t.font.size = Pt(10)
        p_t.font.name = "Segoe UI"
        p_t.font.color.rgb = PALETTE_PPTX["neutral_dark"]
        p_t.space_before = Pt(4)

    def build_slide_3_demographics(self, pyramid_img_path: str, df_age: pd.DataFrame):
        """Slide 3: Pirámide demográfica divergente y tabla de envejecimiento."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_header(slide, "Demografía y Envejecimiento: Alerta en el Relevo Generacional", "Nivel 2: Dimensión Demográfica")

        # Imagen pirámide
        self.add_chart_image(slide, pyramid_img_path, left=Inches(0.8), top=Inches(1.6), width=Inches(7.2))

        # Tabla nativa tramos de edad
        self.add_native_table(slide, df_age, left=Inches(8.3), top=Inches(1.8), width=Inches(4.3), height=Inches(2.4))

        # Cuadro de notas con alerta
        tx = slide.shapes.add_textbox(Inches(8.3), Inches(4.5), Inches(4.3), Inches(2.2))
        tf = tx.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = "Alerta de Relevo Generacional:"
        p_h.font.bold = True
        p_h.font.size = Pt(11)
        p_h.font.color.rgb = PALETTE_PPTX["accent_red"]

        p_t = tf.add_paragraph()
        p_t.text = (
            "• Concentración en edad de jubilación: Las mujeres alcanzan su mayor cuota relativa en el tramo de >=65 años.\n"
            "• En el segmento No ATP, las mujeres mayores de 65 años superan el 56% de representatividad.\n"
            "• El tramo joven (<41 años) solo agrupa a 85 mujeres titulares comerciales en toda Bizkaia."
        )
        p_t.font.size = Pt(10)
        p_t.font.name = "Segoe UI"
        p_t.font.color.rgb = PALETTE_PPTX["neutral_dark"]
        p_t.space_before = Pt(4)

    def build_slide_4_economic_gap(self, dumbbell_img_path: str, kpis: Dict[str, Any]):
        """Slide 4: Brechas de superficie y UTAs (Dumbbell Plot + Tabla de medias)."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_header(slide, "Brechas Económicas y Productivas en Superficie y Mano de Obra", "Nivel 2: Dimensión Productiva")

        # Imagen Dumbbell
        self.add_chart_image(slide, dumbbell_img_path, left=Inches(0.8), top=Inches(1.6), width=Inches(7.2))

        # Tabla nativa de brechas
        atp = kpis.get("atp", {})
        no_atp = kpis.get("no_atp", {})
        df_gaps = pd.DataFrame([
            {"Segmento": "ATP", "Sup H (ha)": atp.get("superficie_media_hombres", 32.6), "Sup M (ha)": atp.get("superficie_media_mujeres", 24.2), "Brecha ha": f"-{atp.get('brecha_superficie_pct', 25.5):.1f}%", "UTAs H": atp.get("utas_media_hombres", 2.1), "UTAs M": atp.get("utas_media_mujeres", 1.4), "Brecha UTAs": f"-{atp.get('brecha_utas_pct', 32.4):.1f}%"},
            {"Segmento": "No ATP", "Sup H (ha)": no_atp.get("superficie_media_hombres", 6.8), "Sup M (ha)": no_atp.get("superficie_media_mujeres", 5.2), "Brecha ha": f"-{no_atp.get('brecha_superficie_pct', 23.5):.1f}%", "UTAs H": no_atp.get("utas_media_hombres", 0.5), "UTAs M": no_atp.get("utas_media_mujeres", 0.4), "Brecha UTAs": f"-{no_atp.get('brecha_utas_pct', 20.0):.1f}%"},
        ])
        self.add_native_table(slide, df_gaps, left=Inches(8.2), top=Inches(1.8), width=Inches(4.4), height=Inches(2.0))

        # Conclusiones
        tx = slide.shapes.add_textbox(Inches(8.2), Inches(4.3), Inches(4.4), Inches(2.4))
        tf = tx.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = "Magnitud de la Brecha Estructural:"
        p_h.font.bold = True
        p_h.font.size = Pt(11)
        p_h.font.color.rgb = PALETTE_PPTX["primary"]

        p_t = tf.add_paragraph()
        p_t.text = (
            "• Brecha de dimensión física: En ATP, las explotaciones de mujeres gestionan 8,3 hectáreas menos de media que las de hombres (-25,5%).\n"
            "• Brecha de capacidad laboral: Las explotaciones femeninas registran 1,42 UTAs frente a 2,10 UTAs masculinas (-32,4%).\n"
            "• Esta doble brecha evidencia menor capitalización inicial y mayores limitaciones para alcanzar economías de escala."
        )
        p_t.font.size = Pt(10)
        p_t.font.name = "Segoe UI"
        p_t.font.color.rgb = PALETTE_PPTX["neutral_dark"]
        p_t.space_before = Pt(4)

    def build_slide_5_subsectors(self, subsectors_img_path: str, df_sub: pd.DataFrame, df_legal: pd.DataFrame):
        """Slide 5: Subsectores OTE UE principales y formas jurídicas."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_header(slide, "Especialización Técnico-Económica (OTE UE) y Personalidad Jurídica", "Nivel 3: Detalle Sectorial")

        # Imagen subsectores
        self.add_chart_image(slide, subsectors_img_path, left=Inches(0.8), top=Inches(1.6), width=Inches(7.2))

        # Tabla nativa formas jurídicas principales
        df_legal_sample = df_legal.head(5)[["Condición Jurídica", "Hombres", "Mujeres", "Total", "% Feminización"]]
        self.add_native_table(slide, df_legal_sample, left=Inches(8.3), top=Inches(1.8), width=Inches(4.3), height=Inches(2.5))

        # Notas
        tx = slide.shapes.add_textbox(Inches(8.3), Inches(4.6), Inches(4.3), Inches(2.2))
        tf = tx.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = "Aspectos Jurídicos y Sectoriales:"
        p_h.font.bold = True
        p_h.font.size = Pt(11)
        p_h.font.color.rgb = PALETTE_PPTX["primary"]

        p_t = tf.add_paragraph()
        p_t.text = (
            "• Predominio de Persona Física: Más del 76% opera como persona física individual.\n"
            "• Titularidad Compartida: Figura emergente con presencia exclusivamente compartida pero aún con baja adopción numérica (7 casos en ATP).\n"
            "• La carne y la horticultura representan los pilares con mayor incorporación de mujeres."
        )
        p_t.font.size = Pt(10)
        p_t.font.name = "Segoe UI"
        p_t.font.color.rgb = PALETTE_PPTX["neutral_dark"]
        p_t.space_before = Pt(4)

    def build_slide_6_territorial(self, territory_img_path: str, df_terr: pd.DataFrame):
        """Slide 6: Ranking territorial comarcal frente a media regional."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.add_header(slide, "Ranking Territorial: Razón de Feminización por Comarcas (ADR)", "Nivel 4: Distribución Geográfica")

        # Imagen ranking territorial
        self.add_chart_image(slide, territory_img_path, left=Inches(0.8), top=Inches(1.6), width=Inches(7.2))

        # Tabla comarcas
        self.add_native_table(slide, df_terr, left=Inches(8.3), top=Inches(1.8), width=Inches(4.3), height=Inches(2.8))

        # Notas finales
        tx = slide.shapes.add_textbox(Inches(8.3), Inches(4.9), Inches(4.3), Inches(2.0))
        tf = tx.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = "Liderazgo Territorial:"
        p_h.font.bold = True
        p_h.font.size = Pt(11)
        p_h.font.color.rgb = PALETTE_PPTX["primary"]

        p_t = tf.add_paragraph()
        p_t.text = (
            "• Enkarterrialde lidera con 72,9 mujeres por 100 hombres en el segmento ATP (42,2% feminización).\n"
            "• Jataondo y Gorbeialde se sitúan próximas al promedio, mientras Urkiola y Lea Artibai registran los índices más masculinizados."
        )
        p_t.font.size = Pt(10)
        p_t.font.name = "Segoe UI"
        p_t.font.color.rgb = PALETTE_PPTX["neutral_dark"]
        p_t.space_before = Pt(4)

    # --------------------------------------------------------------------------
    # ORQUESTADOR COMPLETO
    # --------------------------------------------------------------------------

    def build_full_report(
        self,
        kpis: Dict[str, Any],
        df_age: pd.DataFrame,
        df_sub: pd.DataFrame,
        df_legal: pd.DataFrame,
        df_terr: pd.DataFrame,
        chart_paths: Dict[str, str],
        output_filepath: Union[str, Path] = "exports/presentacion_agraria_bizkaia.pptx",
    ) -> Path:
        """Construye las 6 diapositivas ejecutivas completas y guarda el archivo .pptx."""
        out_path = Path(output_filepath)
        os.makedirs(out_path.parent, exist_ok=True)

        self.build_slide_1_cover(kpis)
        self.build_slide_2_macro_funnel(chart_paths.get("sankey", ""), kpis)
        self.build_slide_3_demographics(chart_paths.get("pyramid", ""), df_age)
        self.build_slide_4_economic_gap(chart_paths.get("dumbbell", ""), kpis)
        self.build_slide_5_subsectors(chart_paths.get("subsectors", ""), df_sub, df_legal)
        self.build_slide_6_territorial(chart_paths.get("territory", ""), df_terr)

        self.prs.save(str(out_path))
        return out_path

    def save(self, filepath: Union[str, Path] = "exports/presentacion_agraria_bizkaia.pptx") -> Path:
        out_path = Path(filepath)
        os.makedirs(out_path.parent, exist_ok=True)
        self.prs.save(str(out_path))
        return out_path

    def to_bytes(self) -> BytesIO:
        buf = BytesIO()
        self.prs.save(buf)
        buf.seek(0)
        return buf


# Alias de compatibilidad previa
PresentationReport = PowerPointReportBuilder


# ==============================================================================
# 4. TEST DE INTEGRACIÓN LOCAL DIRECTO (__main__)
# ==============================================================================

if __name__ == "__main__":
    from src.data_processor import AgrarianDataProcessor
    from src.visualizer import (
        build_age_pyramid_divergent,
        build_dumbbell_economic_gap,
        build_sankey_overview,
        build_subsectors_horizontal_bar,
        build_territorial_ranking,
    )

    print("Iniciando test de integración del motor de exportación (src/exporter.py)...")

    # 1. Cargar datos con AgrarianDataProcessor
    processor = AgrarianDataProcessor.from_file()
    kpis = processor.get_macro_kpis()
    df_age = processor.get_age_distribution(segment="ALL", mode="4_tramos")
    df_sub = processor.get_subsectors_summary(segment="ATP")
    df_legal = processor.get_legal_status_summary(segment="ATP")
    df_terr = processor.get_territorial_summary(segment="ATP")
    df_eco = processor.get_sustainability_summary(segment="ALL")

    # 2. Generar figuras interactivas
    fig_sankey = build_sankey_overview(kpis)
    fig_pyramid = build_age_pyramid_divergent(df_age)
    fig_dumbbell = build_dumbbell_economic_gap(kpis, metric="both")
    fig_subsectors = build_subsectors_horizontal_bar(df_sub, top_n=10)
    fig_territory = build_territorial_ranking(df_terr)

    # 3. Exportar imágenes a exports/charts/ (PNG 300 DPI y SVG)
    chart_files = {
        "sankey": export_chart(fig_sankey, "01_macro_embudo")["png"],
        "pyramid": export_chart(fig_pyramid, "02_piramide_demografica")["png"],
        "dumbbell": export_chart(fig_dumbbell, "03_mancuerna_brechas")["png"],
        "subsectors": export_chart(fig_subsectors, "04_subsectores_ote")["png"],
        "territory": export_chart(fig_territory, "05_ranking_territorial")["png"],
    }

    for key, path_str in chart_files.items():
        p = Path(path_str)
        assert p.exists() and p.stat().st_size > 1000, f"Fallo al generar imagen PNG de {key}"

    # 4. Exportar libro consolidado Excel
    tables_dict = {
        "Demografia 4 Tramos": df_age,
        "Subsectores OTE": df_sub.head(15),
        "Formas Juridicas": df_legal,
        "Ranking Comarcal": df_terr,
        "Sostenibilidad Eco": df_eco,
    }
    excel_path = export_tables_to_excel(tables_dict, output_filepath="exports/tables/tablas_resumen.xlsx")
    assert excel_path.exists() and excel_path.stat().st_size > 1000, "Fallo al generar tablas_resumen.xlsx"

    # 5. Generar presentación PowerPoint completa
    pptx_builder = PowerPointReportBuilder()
    pptx_path = pptx_builder.build_full_report(
        kpis=kpis,
        df_age=df_age,
        df_sub=df_sub,
        df_legal=df_legal,
        df_terr=df_terr,
        chart_paths=chart_files,
        output_filepath="exports/presentacion_agraria_bizkaia.pptx",
    )

    # 6. Comprobar que el archivo se abre y se crea sin corrupciones
    assert pptx_path.exists(), "El archivo PPTX no fue creado."
    assert pptx_path.stat().st_size > 10000, "El tamaño del archivo PPTX es anormalmente pequeño."

    # Validar que python-pptx pueda reabrirlo sin error
    reloaded_prs = Presentation(str(pptx_path))
    assert len(reloaded_prs.slides) == 6, f"Se esperaban 6 diapositivas, se encontraron {len(reloaded_prs.slides)}"

    print("✓ Motor de exportación validado y PowerPoint generado exitosamente.")
